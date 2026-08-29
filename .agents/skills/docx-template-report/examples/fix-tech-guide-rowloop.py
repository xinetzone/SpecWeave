# -*- coding: utf-8 -*-
"""
fix-tech-guide-rowloop.py — 修复 tech-guide-template.docx 表格行循环反模式。

背景：tech-guide-template.docx 的循环表（修订表 + 2/3/4/5 列块表）原先把
纯标签放在数据行单元格内（首格 ``{% for ... %}{{ var }}``、末格
``{{ var }}{% endfor %}``），渲染后单元格横向增生（如 5 列 ×3 数据变成
2 行 ×17 单元格），而渲染示例只做文本断言误判 PASS。

本脚本按 docxtpl 0.20.2 正确机制（三行分离，经 debug-rowloop-patterns.py
G-J 实证）原地修复模板：
  表头行 / for 标记行(渲染移除) / 数据行(仅 {{ 变量 }}) / endfor 标记行(渲染移除)

运行：
  py -3.14 examples/fix-tech-guide-rowloop.py
"""
import os
import re
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(SKILL_DIR, "templates", "tech-guide-template.docx")


def cell_text(tc):
    return "".join(t.text or "" for t in tc.iter(qn("w:t")))


def set_cell_text(tc, text):
    """把单元格首段替换为单 run 文本（保留首 run 的 rPr），清空其余段落。"""
    ps = tc.findall(qn("w:p"))
    p0 = ps[0]
    runs = [r for r in p0.findall(qn("w:r")) if r.find(qn("w:t")) is not None]
    if not runs:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        p0.append(r)
        runs = [r]
    ts = runs[0].findall(qn("w:t"))
    ts[0].text = text
    ts[0].set(qn("xml:space"), "preserve")
    for extra in ts[1:]:
        runs[0].remove(extra)
    for r in runs[1:]:
        p0.remove(r)
    for p in ps[1:]:
        for r in p.findall(qn("w:r")):
            p.remove(r)


def fix_table(tbl_el, for_tag):
    trs = tbl_el.findall(qn("w:tr"))
    header_tr, data_tr = trs[0], trs[1]
    cells = data_tr.findall(qn("w:tc"))

    # 1. 数据行清洗：首格去 for 前缀、末格去 endfor 后缀，只留 {{ 变量 }}
    first_var = re.search(r"(\{\{.*?\}\})\s*$", cell_text(cells[0])).group(1)
    last_var = re.search(r"^\s*(\{\{.*?\}\})", cell_text(cells[-1])).group(1)
    set_cell_text(cells[0], first_var)
    set_cell_text(cells[-1], last_var)

    # 2. 造标记行（deepcopy 数据行保格式，首格放标签，其余格清空）
    def make_marker(tag):
        mtr = deepcopy(data_tr)
        mc = mtr.findall(qn("w:tc"))
        set_cell_text(mc[0], tag)
        for tc in mc[1:]:
            set_cell_text(tc, "")
        return mtr

    for_tr = make_marker(for_tag)
    endfor_tr = make_marker("{%tr endfor %}")
    header_tr.addnext(for_tr)      # 表头 → for 标记行 → 数据行
    data_tr.addnext(endfor_tr)     # 数据行 → endfor 标记行


def main():
    doc = Document(TEMPLATE)
    tbls = doc.tables
    # tbl0=封面品牌栏(1x2)、tbl1=封面信息表(4x2) 无循环；tbl2=修订表；tbl3-6=块表
    fix_table(tbls[2]._tbl, "{%tr for r in revisions %}")
    for t in tbls[3:7]:
        fix_table(t._tbl, "{%tr for row in tbl.rows %}")
    doc.save(TEMPLATE)
    print("[OK] tech-guide-template.docx 行循环已修复为三行分离")

    # 自检：模板物理行数应为 修订表 4 行、块表各 4 行
    chk = Document(TEMPLATE)
    for i, t in enumerate(chk.tables):
        print(f"  tbl{i}: {len(t.rows)}x{len(t.columns)}")


if __name__ == "__main__":
    main()
