# -*- coding: utf-8 -*-
"""隔离测试 2：{%tr if%} 条件行、零行表格、tr-for 行内条件表达式。"""
import os
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docxtpl import DocxTemplate

OUTDIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")
os.makedirs(OUTDIR, exist_ok=True)  # 产物目录不入库，运行时自创建


def mk_p(text=None):
    p = OxmlElement("w:p")
    if text is not None:
        r = OxmlElement("w:r"); t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve"); t.text = text
        r.append(t); p.append(r)
    return p


def cell_set(cell, text):
    p = cell.paragraphs[0]._p
    r = OxmlElement("w:r"); t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve"); t.text = text
    r.append(t); p.append(r)


def try_case(name, build_fn, context, verify=None):
    path = os.path.join(OUTDIR, f"debug2-{name}.docx")
    doc = Document()
    build_fn(doc)
    doc.save(path)
    try:
        tpl = DocxTemplate(path)
        tpl.render(context)
        tpl.save(path)
        chk = Document(path)
        tables = chk.tables
        extra = ""
        if verify:
            extra = verify(chk)
        print(f"[OK] {name} tables={[(len(t.rows), len(t.columns)) for t in tables]} {extra}")
    except Exception as e:
        print(f"[FAIL] {name}: {type(e).__name__}: {str(e)[:200]}")

ctx = {
    "blocks": [
        {"type": "p", "text": "para1"},
        {"type": "table", "cols": 2, "header": ["A", "B"], "rows": [["1", "2"], ["3", "4"]]},
        {"type": "p", "text": "para2"},
    ]
}

def case5(doc):
    """p-for 内含表格（tr-for 行循环），无 p-if —— 整表随循环复制"""
    sect = doc.element.body.find(qn("w:sectPr"))
    sect.addprevious(mk_p('{%p for blk in blocks %}'))
    tbl = doc.add_table(rows=2, cols=2)
    cell_set(tbl.rows[0].cells[0], "{{ blk.header[0] }}")
    cell_set(tbl.rows[0].cells[1], "{{ blk.header[1] }}")
    cell_set(tbl.rows[1].cells[0], "{%tr for row in blk.rows %}{{ row[0] }}")
    cell_set(tbl.rows[1].cells[1], "{{ row[1] }}{%tr endfor %}")
    sect.addprevious(mk_p('{%p endfor %}'))

def case6(doc):
    """p-for 内含表格；tr-if 条件行（表头/数据行均条件化），不匹配时零行"""
    sect = doc.element.body.find(qn("w:sectPr"))
    sect.addprevious(mk_p('{%p for blk in blocks %}'))
    tbl = doc.add_table(rows=2, cols=2)
    cond = "blk.type == 'table' and blk.cols == 2"
    cell_set(tbl.rows[0].cells[0], f"{{%tr if {cond} %}}{{{{ blk.header[0] }}}}")
    cell_set(tbl.rows[0].cells[1], "{{ blk.header[1] }}{%tr endif %}")
    cell_set(tbl.rows[1].cells[0],
             f"{{%tr for row in (blk.rows if {cond} else []) %}}{{{{ row[0] }}}}")
    cell_set(tbl.rows[1].cells[1], "{{ row[1] }}{%tr endfor %}")
    sect.addprevious(mk_p('{%p endfor %}'))

def case7(doc):
    """case6 + 一个 p-if 段落（验证 tr-if 与 p-if 共存）"""
    sect = doc.element.body.find(qn("w:sectPr"))
    sect.addprevious(mk_p('{%p for blk in blocks %}'))
    sect.addprevious(mk_p('{%p if blk.type == "p" %}'))
    sect.addprevious(mk_p("{{ blk.text }}"))
    sect.addprevious(mk_p('{%p endif %}'))
    tbl = doc.add_table(rows=2, cols=2)
    cond = "blk.type == 'table' and blk.cols == 2"
    cell_set(tbl.rows[0].cells[0], f"{{%tr if {cond} %}}{{{{ blk.header[0] }}}}")
    cell_set(tbl.rows[0].cells[1], "{{ blk.header[1] }}{%tr endif %}")
    cell_set(tbl.rows[1].cells[0],
             f"{{%tr for row in (blk.rows if {cond} else []) %}}{{{{ row[0] }}}}")
    cell_set(tbl.rows[1].cells[1], "{{ row[1] }}{%tr endfor %}")
    sect.addprevious(mk_p('{%p endfor %}'))

def verify_text(chk):
    txt = "\n".join(c.text for t in chk.tables for r in t.rows for c in r.cells)
    return f"含'1,2'数据={'1' in txt and '3' in txt} 空表数={sum(1 for t in chk.tables if len(t.rows)==0)}"

try_case("case5-table-in-pfor-noif", case5, ctx)
try_case("case6-trif-zero-rows", case6, ctx, verify=verify_text)
try_case("case7-trif-with-pif", case7, ctx, verify=verify_text)
