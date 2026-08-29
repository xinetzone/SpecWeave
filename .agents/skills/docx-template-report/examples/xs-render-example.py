# -*- coding: utf-8 -*-
"""
xs-render-example.py — xs-sdk-guide-template.docx 渲染验证示例。

覆盖全部块类型：封面字段 / 修订记录行循环(3行) / 目录域 / h1-h4 / 正文 / 列表 /
代码块(多行) / shell 行 / 警告行 / 2-3-4-5 列表格 / 显式分页。
渲染后执行 37 项断言，验证样式、编号链、页眉、表格保真与品牌脱敏。
演示数据全部为虚构中立数据，渲染产物不含任何品牌/人员关键词。

运行：
  py -3.14 examples/xs-render-example.py
"""
import os
import zipfile
from docx import Document
from docx.oxml.ns import qn
from docxtpl import DocxTemplate

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(BASE, "templates", "xs-sdk-guide-template.docx")
OUT = os.path.join(BASE, "examples", "output", "xs-template-example.docx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)  # 产物目录不入库，运行时自创建

# 品牌检测词表以扫描器 scan-brand-residue.py 的 KEYWORDS 为唯一事实源，
# 本示例不另立词表（检测签名单点维护，随扫描器更新）
import importlib.util
_spec = importlib.util.spec_from_file_location(
    "scan_brand_residue",
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "scan-brand-residue.py"))
_scan = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_scan)
BRAND_KEYWORDS = _scan.KEYWORDS

# 演示数据均为虚构中立数据，不含任何真实品牌/人员信息
context = {
    "doc_title": "XS SDK 用户使用指南",
    "doc_status": "正式发布",
    "doc_version": "1.1.0",
    "doc_author": "张三",
    "doc_date": "2025-10-15",
    "doc_reviewer": "李四",
    "company": "示例科技有限公司",
    "copyright_notice": "（版本所有，翻版必究）",
    "revisions": [
        {"version": "1.0.0", "author": "王五", "date": "2025-07-02",
         "description": "初始版本", "reviewer": "李四"},
        {"version": "1.1.0", "author": "王五", "date": "2025-10-15",
         "description": "1.工具链移除对 caffe 环境的依赖；2.新增算子支持：matmul, softmax",
         "reviewer": "李四"},
        {"version": "1.2.0", "author": "测试员", "date": "2026-08-29",
         "description": "模板渲染验证追加行", "reviewer": "审核员"},
    ],
    "blocks": [
        {"type": "h1", "text": "开发环境准备"},
        {"type": "p", "text": "本工具链当前仅支持在 Linux 操作系统环境下运行，资源准备请先解压 SDK 压缩包。"},
        {"type": "list", "text": "Conda 虚拟环境配置"},
        {"type": "list", "text": "Docker 容器环境配置"},
        {"type": "h2", "text": "Conda 环境配置"},
        {"type": "h3", "text": "构建 xsenv 环境"},
        {"type": "h4", "text": "环境变量说明"},
        {"type": "p", "text": "创建并激活 Python 环境，依次执行以下命令："},
        {"type": "code", "lines": [
            "cd release",
            "conda env create --file=xs.yaml",
            "conda activate xsenv",
            "pip install xs_sdk-1.1.0-py3-none-any.whl",
        ]},
        {"type": "shell", "text": "$ conda activate xsenv"},
        {"type": "shell", "text": "$ docker run -it --rm -v xxx/npuusertools:/home/ xs:latest"},
        {"type": "warn", "text": "注意：pytorch 需要使用 torch.jit.save 保存的模型。"},
        {"type": "table", "cols": 2,
         "header": ["参数", "说明"],
         "rows": [["-i", "以交互模式运行容器"], ["-t", "为容器重新分配一个伪终端"]]},
        {"type": "table", "cols": 3,
         "header": ["输入类型", "支持格式", "备注"],
         "rows": [["图像", "jpg/png", "单输入"], ["视频", "mp4/avi", "批量帧"]]},
        {"type": "table", "cols": 4,
         "header": ["参数名", "类型", "默认值", "说明"],
         "rows": [["batch_size", "int", "1", "批大小"], ["device", "str", "npu", "运行设备"]]},
        {"type": "table", "cols": 5,
         "header": ["参数名称", "必填", "类型", "默认值", "说明"],
         "rows": [["model", "是", "str", "-", "模型路径"], ["output", "否", "str", "out", "输出目录"]]},
        {"type": "pagebreak"},
        {"type": "h1", "text": "量化原理"},
        {"type": "p", "text": "量化通过降低权重与激活的位宽以提升 NPU 推理效率。"},
    ],
}

# ---------------------------------------------------------------- 渲染

doc = DocxTemplate(TEMPLATE)
doc.render(context)
doc.save(OUT)
print(f"[OK] 渲染输出: {OUT}")

# ---------------------------------------------------------------- 断言

chk = Document(OUT)
# 注意：doc.paragraphs 只含 body 直接子段落，不覆盖 w:sdt 块（封面）内段落；
# 用 body 全树 w:t 迭代，确保封面 sdt 文本纳入断言。
body_text = "\n".join(t.text or "" for t in chk.element.body.iter(qn("w:t")))
table_text = "\n".join(
    cell.text for tbl in chk.tables for row in tbl.rows for cell in row.cells
)
passed = []

def check(name, cond):
    assert cond, f"断言失败: {name}"
    passed.append(name)

# A. 文件与封面
check("A1 输出文件非空", os.path.getsize(OUT) > 20_000)
check("A2 封面标题渲染", "XS SDK 用户使用指南" in body_text)
check("A3 版本/状态渲染", "1.1.0" in body_text and "正式发布" in body_text)
check("A4 无残留 Jinja 标签", "{%" not in body_text and "{{" not in body_text and "{%" not in table_text and "{{" not in table_text)
# 封面 sdt 内 作者/完成日期/审核 三字段
cover_sdt = next(ch for ch in chk.element.body.iter(qn("w:sdt")))
cover_text = "\n".join(t.text or "" for t in cover_sdt.iter(qn("w:t")))
check("A5 封面作者/日期/审核渲染", "张三" in cover_text and "2025-10-15" in cover_text and "李四" in cover_text)
# 模板内三字段确为变量标签（非硬编码）
with zipfile.ZipFile(TEMPLATE) as z:
    tpl_xml = z.read("word/document.xml").decode("utf-8")
check("A6 模板封面三字段已变量化", all(
    tag in tpl_xml for tag in ("{{ doc_author }}", "{{ doc_date }}", "{{ doc_reviewer }}")))

# B. 封面公司名（隐形表）与页眉
check("B1 封面公司名", "示例科技有限公司" in table_text)
hdr_text = ""
for sec in chk.sections:
    for hdr in [sec.header, sec.first_page_header]:
        for p in hdr.paragraphs:
            hdr_text += p.text + "\n"
        for tbl in hdr.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    hdr_text += cell.text + "\n"
check("B2 页眉公司名标签已渲染", "示例科技有限公司" in hdr_text)
# B3：页眉脱敏——VML 水印/品牌图形必须全部移除（logo 已占位化）
with zipfile.ZipFile(OUT) as z:
    hdr_xml = "".join(
        z.read(n).decode("utf-8", "ignore")
        for n in z.namelist() if n.startswith("word/header")
    )
check("B3 页眉零水印零图形",
      "<w:pict" not in hdr_xml and "<w:drawing" not in hdr_xml
      and all(kw not in hdr_xml for kw in BRAND_KEYWORDS))

# C. 更新记录页
check("C1 更新记录标题存在", "更新记录" in body_text)
rev_tbl = chk.tables[1]
check("C2 修订表 3 数据行+1 表头=4 行", len(rev_tbl.rows) == 4)
check("C3 修订表循环数据", "模板渲染验证追加行" in table_text and "2025-10-15" in table_text)
# 更新记录 H1 的 numId=0 抑制编号
h1_upd = next(p for p in chk.paragraphs if p.text.strip() == "更新记录")
numpr = h1_upd._p.find(qn("w:pPr")).find(qn("w:numPr"))
numid_val = numpr.find(qn("w:numId")).get(qn("w:val")) if numpr is not None else None
check("C4 更新记录 numId=0 抑制编号", numid_val == "0")
# 修订表深色底纹
rev_hdr_shd = rev_tbl.rows[0].cells[0]._tc.find(qn("w:tcPr")).find(qn("w:shd"))
check("C5 修订表头 191919 黑底", rev_hdr_shd is not None and rev_hdr_shd.get(qn("w:fill")) == "191919")

# D. 目录域 + G. 部件保真（合并一次 zip 读取）
with zipfile.ZipFile(OUT) as z:
    doc_xml = z.read("word/document.xml").decode("utf-8", "ignore")
    settings_xml = z.read("word/settings.xml").decode("utf-8")
    numbering_xml = z.read("word/numbering.xml").decode("utf-8")
    names = z.namelist()
check("D1 目录标题存在", "目录" in body_text)
check("D2 TOC 域存在", "TOC " in doc_xml and "fldChar" in doc_xml)
check("D3 updateFields 已开启", "updateFields" in settings_xml)

# E. 正文块
h1s = [p for p in chk.paragraphs if p.style and p.style.style_id == "2"]
h2s = [p for p in chk.paragraphs if p.style and p.style.style_id == "3"]
check("E1 H1 数量=3(更新记录+2块)", len(h1s) == 3)
check("E2 H2 数量=1", len(h2s) == 1)
code_ps = [p for p in chk.paragraphs if p.style and p.style.style_id == "64"]
# 4 code 行 + 2 shell 行 = 6
check("E3 代码块行数=6(4code+2shell)", len(code_ps) == 6)
# shell 行颜色 215F9A
shell_p = next(p for p in code_ps if p.text.startswith("$ conda"))
shell_color = None
for r in shell_p._p.findall(qn("w:r")):
    rpr = r.find(qn("w:rPr"))
    if rpr is not None and rpr.find(qn("w:color")) is not None:
        shell_color = rpr.find(qn("w:color")).get(qn("w:val"))
check("E4 shell 行 215F9A 蓝色", shell_color == "215F9A")
# warn 行颜色 EE0000
warn_p = next(p for p in chk.paragraphs if "torch.jit.save" in p.text)
warn_color = None
for r in warn_p._p.findall(qn("w:r")):
    rpr = r.find(qn("w:rPr"))
    if rpr is not None and rpr.find(qn("w:color")) is not None:
        warn_color = rpr.find(qn("w:color")).get(qn("w:val"))
check("E5 警告行 EE0000 红色", warn_color == "EE0000")
# 列表项样式 46
list_ps = [p for p in chk.paragraphs if p.style and p.style.style_id == "46"]
check("E6 列表项 2 个(样式46)", len(list_ps) == 2)

# F. 表格（4 个块表 + 封面/修订表 = 6）
check("F1 表格总数=6", len(chk.tables) == 6)
blk_tables = chk.tables[2:]
dims = [(len(t.rows), len(t.columns)) for t in blk_tables]
# 2列: 1表头+2数据=3行；3列: 3行；4列: 3行；5列: 3行
check("F2 块表维度 3x2/3x3/3x4/3x5", dims == [(3, 2), (3, 3), (3, 4), (3, 5)], )
check("F3 表格数据渲染", "batch_size" in table_text and "伪终端" in table_text and "模型路径" in table_text)
# 表头黑底白字
t_hdr_shd = blk_tables[2].rows[0].cells[0]._tc.find(qn("w:tcPr")).find(qn("w:shd"))
check("F4 块表表头黑底 000000", t_hdr_shd.get(qn("w:fill")) == "000000")

# G. 样式与编号链保真
check("G1 code 样式(64)存在", any(
    s.get(qn("w:styleId")) == "64" for s in chk.styles.element.findall(qn("w:style"))))
check("G2 numbering.xml 含标题多级列表绑定", 'w:pStyle w:val="2"' in numbering_xml and 'w:pStyle w:val="3"' in numbering_xml)
check("G3 theme1.xml 保留", "word/theme/theme1.xml" in names)
check("G4 渲染产物零媒体(脱敏)", not any(n.startswith("word/media/") for n in names))

# H. 品牌资产脱敏专项（模板与渲染产物双查）
with zipfile.ZipFile(TEMPLATE) as z:
    tpl_all = "".join(
        z.read(n).decode("utf-8", "ignore")
        for n in z.namelist() if n.endswith(".xml")
    )
    tpl_names = z.namelist()
    core_xml = z.read("docProps/core.xml").decode("utf-8", "ignore")
check("H1 模板零品牌词零水印元素",
      all(kw not in tpl_all for kw in BRAND_KEYWORDS) and "<w:pict" not in tpl_all)
check("H2 模板零媒体文件", not any(n.startswith("word/media/") for n in tpl_names))
check("H3 logo 占位标签存在", "{{ header_logo }}" in tpl_all and "{{ cover_logo }}" in tpl_all)
check("H4 docProps 元数据已脱敏", all(kw not in core_xml for kw in BRAND_KEYWORDS))
with zipfile.ZipFile(OUT) as z:
    out_all = "".join(
        z.read(n).decode("utf-8", "ignore")
        for n in z.namelist() if n.endswith(".xml")
    )
    out_names = z.namelist()
check("H5 渲染产物零品牌词零水印",
      all(kw not in out_all for kw in BRAND_KEYWORDS) and "<w:pict" not in out_all
      and not any(n.startswith("word/media/") for n in out_names))
check("H6 页眉无残留 Jinja 标签", "{{" not in hdr_xml and "{%" not in hdr_xml)

print(f"\n[PASS] 全部 {len(passed)} 项断言通过：")
for i, name in enumerate(passed, 1):
    print(f"  {i:2d}. {name}")
