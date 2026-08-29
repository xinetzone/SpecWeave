# -*- coding: utf-8 -*-
"""补充分析 2：两个 sdt 块内容（封面标题块/目录块）、T2 黑底表头、分页空段。"""
from docx import Document
from docx.oxml.ns import qn
import sys
import glob

# 源文档在本地 .chaos 工作区（不入库）；以通配符定位，不固化源文件名
SRC = sys.argv[1] if len(sys.argv) > 1 else (
    glob.glob(r"d:\AI\.chaos\tests\old\work\doc\*SDK*指南*.docx") or [None])[0]
doc = Document(SRC)
body = doc.element.body

def el(parent, tag):
    return parent.find(qn(tag))

def p_text(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t")))

def p_style(p):
    ppr = el(p, "w:pPr")
    if ppr is None:
        return None
    ps = el(ppr, "w:pStyle")
    return ps.get(qn("w:val")) if ps is not None else None

children = list(body)
sdts = [ch for ch in children if ch.tag == qn("w:sdt")]
print(f"sdt 块总数: {len(sdts)}")
for si, sdt in enumerate(sdts):
    print(f"\n=== sdt[{si}] ===")
    paras = list(sdt.iter(qn("w:p")))
    print(f"  段落数: {len(paras)}")
    has_drawing = len(list(sdt.iter(qn("w:drawing")))) > 0
    has_toc_field = "TOC" in "".join(t.text or "" for t in sdt.iter(qn("w:instrText")))
    instr = [t.text for t in sdt.iter(qn("w:instrText")) if t.text and t.text.strip()]
    print(f"  含图片: {has_drawing}  含TOC域: {has_toc_field}  域指令: {instr[:3]}")
    for p in paras[:8]:
        txt = p_text(p)
        if txt.strip() or p is paras[0]:
            # run rPr
            r0 = p.findall(qn("w:r"))
            rpr_desc = ""
            if r0:
                rpr = el(r0[0], "w:rPr")
                if rpr is not None:
                    sz = el(rpr, "w:sz")
                    rf = el(rpr, "w:rFonts")
                    b = el(rpr, "w:b")
                    color = el(rpr, "w:color")
                    rpr_desc = f"sz={sz.get(qn('w:val')) if sz is not None else '-'} ea={rf.get(qn('w:eastAsia')) if rf is not None else '-'} b={'Y' if b is not None else 'N'} color={color.get(qn('w:val')) if color is not None else '-'}"
            jc = el(el(p, "w:pPr") or p, "w:jc")
            jcv = jc.get(qn("w:val")) if jc is not None else "-"
            print(f"    style={p_style(p)} jc={jcv} [{rpr_desc}] text={txt[:50]!r}")

print("\n=== T2 (6x2, style=26) 表头行完整格式 ===")
t2 = doc.tables[2]
for ri, row in enumerate(t2.rows[:2]):
    for ci, cell in enumerate(row.cells):
        tcpr = cell._tc.find(qn("w:tcPr"))
        shd = el(tcpr, "w:shd") if tcpr is not None else None
        shd_fill = shd.get(qn("w:fill")) if shd is not None else None
        for p in cell.paragraphs:
            for r in p._p.findall(qn("w:r")):
                rpr = el(r, "w:rPr")
                if rpr is not None:
                    color = el(rpr, "w:color")
                    b = el(rpr, "w:b")
                    sz = el(rpr, "w:sz")
                    print(f"  r{ri}c{ci}: shd={shd_fill} color={color.get(qn('w:val')) if color is not None else '-'} b={'Y' if b is not None else 'N'} sz={sz.get(qn('w:val')) if sz is not None else '-'} text={p_text(p._p)[:30]!r}")

print("\n=== 分页空段（body 子元素 idx 4,5,7 附近）===")
for i in [4, 5, 6, 7]:
    ch = children[i]
    if ch.tag == qn("w:p"):
        brs = [br.get(qn("w:type")) for br in ch.iter(qn("w:br"))]
        print(f"  child[{i}] p style={p_style(ch)} brs={brs} text={p_text(ch)[:20]!r}")

print("\n=== 封面 sdt[0] 之后/封面表 T0 的前后顺序确认 ===")
for i in range(0, 9):
    ch = children[i]
    tag = ch.tag.split("}")[1]
    if tag == "p":
        print(f"  [{i}] p style={p_style(ch)} text={p_text(ch)[:40]!r}")
    elif tag == "tbl":
        ft = next(ch.iter(qn("w:t")), None)
        print(f"  [{i}] tbl first={ft.text if ft is not None else ''!r} rows={len(ch.findall(qn('w:tr')))}")
    else:
        print(f"  [{i}] {tag}")

print("\n=== 正文样式26/81 表格中表头黑底行统计 ===")
for ti, tbl in enumerate(doc.tables):
    if tbl.style and tbl.style.style_id in ("26", "81"):
        hdr = tbl.rows[0]
        fills = []
        for cell in hdr.cells:
            tcpr = cell._tc.find(qn("w:tcPr"))
            s = el(tcpr, "w:shd") if tcpr is not None else None
            fills.append(s.get(qn("w:fill")) if s is not None else None)
        print(f"  T{ti} style={tbl.style.style_id} {len(tbl.rows)}x{len(tbl.columns)} 表头fills={fills} 首格={hdr.cells[0].text[:15]!r}")
