# -*- coding: utf-8 -*-
"""隔离测试 docxtpl 标签组合：p-for 内嵌套表格 tr-for、嵌套 p-for。"""
import os
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docxtpl import DocxTemplate

OUTDIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")
os.makedirs(OUTDIR, exist_ok=True)


def mk_p(tag_text=None, style=None):
    p = OxmlElement("w:p")
    if style:
        ppr = OxmlElement("w:pPr")
        ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), style); ppr.append(ps)
        p.append(ppr)
    if tag_text:
        r = OxmlElement("w:r"); t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve"); t.text = tag_text
        r.append(t); p.append(r)
    return p


def add_table(doc, anchor, ncols):
    tbl = doc.add_table(rows=2, cols=ncols)
    for i, cell in enumerate(tbl.rows[0].cells):
        cell.paragraphs[0]._p.append(mk_p_run(f"{{{{ blk.header[{i}] }}}}"))
    for i, cell in enumerate(tbl.rows[1].cells):
        txt = f'{{%tr for row in blk.rows %}}{{{{ row[{i}] }}}}' if i == 0 else (
            f'{{{{ row[{i}] }}}}{{%tr endfor %}}' if i == ncols - 1 else f'{{{{ row[{i}] }}}}')
        cell.paragraphs[0]._p.append(mk_p_run(txt))
    return tbl


def mk_p_run(text):
    r = OxmlElement("w:r"); t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve"); t.text = text
    r.append(t)
    return r


def try_case(name, build_fn, context):
    path = os.path.join(OUTDIR, f"debug-{name}.docx")
    doc = Document()
    build_fn(doc)
    doc.save(path)
    try:
        tpl = DocxTemplate(path)
        tpl.render(context)
        tpl.save(path)
        print(f"[OK] {name}")
        return True
    except Exception as e:
        print(f"[FAIL] {name}: {type(e).__name__}: {e}")
        return False

ctx = {
    "blocks": [
        {"type": "p", "text": "hello"},
        {"type": "table", "cols": 2, "header": ["a", "b"], "rows": [["1", "2"], ["3", "4"]]},
        {"type": "code", "lines": ["l1", "l2", "l3"]},
    ]
}

# 用基底模板（带 Heading 样式）——直接用 python-docx 默认模板即可（测标签逻辑）
def case1(doc):
    """p-for + p-if 段落（无表格）"""
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    for tag in ['{%p for blk in blocks %}',
                '{%p if blk.type == "p" %}']:
        p = mk_p(tag); sect.addprevious(p)
    p = mk_p("{{ blk.text }}"); sect.addprevious(p)
    p = mk_p("{%p endif %}"); sect.addprevious(p)
    p = mk_p("{%p endfor %}"); sect.addprevious(p)

def case2(doc):
    """p-for + p-if 内含表格 tr-for"""
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    sect.addprevious(mk_p('{%p for blk in blocks %}'))
    sect.addprevious(mk_p('{%p if blk.type == "table" and blk.cols == 2 %}'))
    add_table(doc, sect, 2)
    sect.addprevious(mk_p('{%p endif %}'))
    sect.addprevious(mk_p('{%p endfor %}'))

def case3(doc):
    """p-for + p-if 内含嵌套 p-for（code lines）"""
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    sect.addprevious(mk_p('{%p for blk in blocks %}'))
    sect.addprevious(mk_p('{%p if blk.type == "code" %}'))
    sect.addprevious(mk_p('{%p for line in blk.lines %}'))
    sect.addprevious(mk_p("{{ line }}"))
    sect.addprevious(mk_p('{%p endfor %}'))
    sect.addprevious(mk_p('{%p endif %}'))
    sect.addprevious(mk_p('{%p endfor %}'))

def case4(doc):
    """p-for + p-if 段落 + p-if 表格（case1+case2 合并，贴近真实模板）"""
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    sect.addprevious(mk_p('{%p for blk in blocks %}'))
    sect.addprevious(mk_p('{%p if blk.type == "p" %}'))
    sect.addprevious(mk_p("{{ blk.text }}"))
    sect.addprevious(mk_p('{%p endif %}'))
    sect.addprevious(mk_p('{%p if blk.type == "table" and blk.cols == 2 %}'))
    add_table(doc, sect, 2)
    sect.addprevious(mk_p('{%p endif %}'))
    sect.addprevious(mk_p('{%p endfor %}'))

try_case("case1-pif", case1, ctx)
try_case("case2-tr-in-pfor", case2, ctx)
try_case("case3-nested-pfor", case3, ctx)
try_case("case4-combined", case4, ctx)
