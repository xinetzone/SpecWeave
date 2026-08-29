# -*- coding: utf-8 -*-
"""第三轮补充：封面表格内容 / 页眉页脚XML直查 / 特殊颜色段落定位"""
import sys
from docx import Document
from docx.oxml.ns import qn

doc = Document(sys.argv[1] if len(sys.argv) > 1 else
               r"d:\AI\.chaos\tests\old\work\doc\XMNN_SDK_使用指南v1.1.0.docx")

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

print("=== 页眉页脚（直接遍历 XML）===")
for idx, s in enumerate(doc.sections):
    for name, obj in (("header", s.header), ("footer", s.footer),
                      ("first_page_header", s.first_page_header),
                      ("first_page_footer", s.first_page_footer)):
        try:
            if obj.is_linked_to_previous:
                print(f"[sec{idx}] {name}: linked to previous")
                continue
            print(f"\n[sec{idx}] {name}:")
            for p in obj.paragraphs:
                if p.text.strip() or p._p.findall(f".//{{{W}}}r"):
                    fields = [t.text for t in p._p.findall(f".//{{{W}}}instrText")]
                    print(f"  para style={p.style.name!r} align={p.alignment} "
                          f"text={p.text!r} fields={fields}")
            for ti, t in enumerate(obj.tables):
                print(f"  table#{ti}: {len(t.rows)}行x{len(t.columns)}列")
                for ri, row in enumerate(t.rows):
                    cells = [c.text.strip()[:40] for c in row.cells]
                    print(f"    row{ri}: {cells}")
                    for ci, c in enumerate(row.cells):
                        for p in c.paragraphs:
                            for run in p.runs:
                                if run.text.strip():
                                    f = run.font
                                    print(f"      cell{ci} run: {run.text[:30]!r} "
                                          f"size={f.size} bold={f.bold} color={f.color.rgb if f.color and f.color.type else None}")
                        # 图片
                        if c._tc.findall(f".//{{{W}}}drawing") or c._tc.findall(f".//{{{W}}}pict"):
                            print(f"      cell{ci}: [含图片/绘图]")
        except Exception as e:
            print(f"[sec{idx}] {name}: ERROR {e}")

print("\n\n=== 正文前 2 个表格（T0 封面 / T1 修订记录）单元格全文 ===")
for ti, t in enumerate(doc.tables[:2]):
    print(f"\n--- T{ti}: {len(t.rows)}行x{len(t.columns)}列 style={t.style.name} ---")
    for ri, row in enumerate(t.rows):
        for ci, c in enumerate(row.cells):
            txt = c.text.strip()
            if txt:
                print(f"  [r{ri}c{ci}] {txt[:120]!r}")
            for p in c.paragraphs:
                for run in p.runs:
                    if run.text.strip():
                        f = run.font
                        ea = run._r.find(f".//{{{W}}}rFonts")
                        ea_font = ea.get(qn("w:eastAsia")) if ea is not None else None
                        ascii_font = ea.get(qn("w:ascii")) if ea is not None else None
                        print(f"      run: {run.text[:50]!r} ascii={ascii_font} eastAsia={ea_font} "
                              f"size={f.size.pt if f.size else None}pt bold={f.bold} "
                              f"color={f.color.rgb if f.color and f.color.type else None}")
            if c._tc.findall(f".//{{{W}}}drawing"):
                print(f"  [r{ri}c{ci}] [含图片]")

print("\n\n=== 含 #4E95D9 / #EE0000 / 15pt 的段落全文 ===")
for p in doc.paragraphs:
    for run in p.runs:
        col = None
        if run.font.color and run.font.color.type:
            col = str(run.font.color.rgb)
        sz = run.font.size.pt if run.font.size else None
        if col in ("4E95D9", "EE0000") or sz == 15.0:
            print(f"  style={p.style.name!r} color={col} size={sz} bold={run.bold} "
                  f"text={run.text[:80]!r}")
            break

print("\n\n=== 正文图片所在段落（drawing 定位）===")
for i, p in enumerate(doc.paragraphs):
    drawings = p._p.findall(f".//{{{W}}}drawing")
    if drawings:
        # 图片前后的段落文字
        ctx_before = doc.paragraphs[i-1].text[:50] if i > 0 else ""
        print(f"  para[{i}] style={p.style.name!r} 图片x{len(drawings)} 前文={ctx_before!r}")
